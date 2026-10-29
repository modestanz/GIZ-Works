# -*- coding: utf-8 -*-
"""Apply the V5 revisions to the F1-D4 V4 Word report, editing it in place.

Usage: python3 build_v5.py <V4.docx> <out.docx> [pages.json]
pages.json (optional) maps heading text -> page number, produced by toc_pages.py.
"""
import copy, json, re, sys
from lxml import etree
from docx import Document
import content as C

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
q = lambda t: "{%s}%s" % (W, t)
src, out = sys.argv[1], sys.argv[2]
pages = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}

doc = Document(src)
body = doc.element.body

# ------------------------------------------------------------------ helpers
def el(tag, **attrs):
    e = etree.Element(q(tag))
    for k, v in attrs.items():
        e.set(q(k), str(v))
    return e

def sub(parent, tag, **attrs):
    e = el(tag, **attrs); parent.append(e); return e

def text_of(e):
    return "".join(t.text or "" for t in e.iter(q("t")))

def style_of(p):
    s = p.find(q("pPr") + "/" + q("pStyle"))
    return s.get(q("val")) if s is not None else ""

def top_paras():
    return [k for k in body if k.tag == q("p")]

def find_p(pred, start_after=None):
    seen = start_after is None
    for k in body:
        if not seen:
            seen = k is start_after
            continue
        if k.tag == q("p") and pred(k):
            return k
    raise KeyError("paragraph not found")

def startswith(s):
    return lambda p: text_of(p).strip().startswith(s)

def heading(s, lvl=None):
    def f(p):
        st = style_of(p)
        ok = st.startswith("Heading") if lvl is None else st == "Heading%d" % lvl
        return ok and s in text_of(p)
    return f

def rpr(bold=False, italic=False, sz=None, hl=False):
    r = el("rPr")
    sub(r, "rFonts", ascii="Times New Roman", hAnsi="Times New Roman", cs="Times New Roman")
    if bold: sub(r, "b"); sub(r, "bCs")
    if italic: sub(r, "i"); sub(r, "iCs")
    sub(r, "color", val="000000", themeColor="text1")
    if sz: sub(r, "sz", val=sz); sub(r, "szCs", val=sz)
    if hl: sub(r, "highlight", val="yellow")
    return r

def add_runs(p, text, bold=False, italic=False, sz=None):
    for tok in re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text):
        if not tok:
            continue
        b, hl = bold, False
        if tok.startswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("[["):
            tok, hl = "[" + tok[2:-2] + "]", True
        r = sub(p, "r"); r.append(rpr(b, italic, sz, hl))
        t = sub(r, "t"); t.text = tok
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

def ppr(style=None, jc=None, spacing=None, num=None, keep_next=False, page_break=False):
    pp = el("pPr")
    if style: sub(pp, "pStyle", val=style)
    if keep_next: sub(pp, "keepNext")
    if page_break: sub(pp, "pageBreakBefore")
    if num:
        n = sub(pp, "numPr"); sub(n, "ilvl", val=num[1]); sub(n, "numId", val=num[0])
    if spacing:
        sub(pp, "spacing", **spacing)
    if jc: sub(pp, "jc", val=jc)
    pp.append(rpr())
    return pp

BM = [900000]
def bookmark_wrap(p, name=None):
    BM[0] += 1
    name = name or "_Toc9%07d" % BM[0]
    bs = el("bookmarkStart", id=BM[0], name=name)
    p.insert(1, bs)
    p.append(el("bookmarkEnd", id=BM[0]))
    return name

def P(kind, text=""):
    p = el("p")
    if kind == "p":
        p.append(ppr(jc="both")); add_runs(p, text)
    elif kind == "lead":
        p.append(ppr(jc="both"))
        m = re.match(r"(In brief:)(.*)", text, re.S)
        add_runs(p, m.group(1), bold=True, italic=True); add_runs(p, m.group(2), italic=True)
    elif kind == "note":
        p.append(ppr(jc="both")); add_runs(p, text, italic=True, sz=20)
    elif kind == "b":
        p.append(ppr(style="ListParagraph", jc="both", num=(6, 0))); add_runs(p, text)
    elif kind == "h1":
        p.append(ppr(style="Heading1"))
        r = sub(p, "r"); r.append(rpr()); sub(r, "br", type="page")
        add_runs(p, text, bold=True, sz=28); bookmark_wrap(p)
    elif kind == "h2":
        p.append(ppr(style="Heading2", spacing={"after": 120})); add_runs(p, text, bold=True, sz=24); bookmark_wrap(p)
    elif kind == "h3":
        p.append(ppr(style="Heading3")); add_runs(p, text, bold=True, italic=True, sz=24)
    elif kind == "blank":
        p.append(ppr())
    return p

TEMPLATE_TBL = doc.tables[1]._tbl
CAPTIONS = {}  # key -> caption paragraph

def caption(key, text):
    p = el("p"); p.append(ppr(keep_next=True, spacing={"before": 160, "after": 80}))
    add_runs(p, "Table 0: ", bold=True, sz=20); add_runs(p, text, italic=True, sz=20)
    CAPTIONS[key] = p
    return p

def cell(text, w, header=False):
    tc = el("tc"); tp = sub(tc, "tcPr"); sub(tp, "tcW", w=w, type="dxa")
    for i, line in enumerate(text.split("\n")):
        p = sub(tc, "p"); pp = sub(p, "pPr"); sub(pp, "spacing", line=252, lineRule="auto"); pp.append(rpr())
        add_runs(p, line, bold=header, sz=18)
    return tc

def table(key, cap, headers, rows, widths):
    assert sum(widths) == 9026, (key, sum(widths))
    tbl = el("tbl")
    tbl.append(copy.deepcopy(TEMPLATE_TBL.find(q("tblPr"))))
    g = sub(tbl, "tblGrid")
    for w in widths: sub(g, "gridCol", w=w)
    tr = sub(tbl, "tr"); trp = sub(tr, "trPr"); sub(trp, "tblHeader")
    for h, w in zip(headers, widths): tr.append(cell(h, w, True))
    for row in rows:
        tr = sub(tbl, "tr"); trp = sub(tr, "trPr"); sub(trp, "cantSplit")
        for c, w in zip(row, widths): tr.append(cell(c, w))
    return ([caption(key, cap)] if cap else []) + [tbl, P("blank")]

def render(blocks, annex=False):
    out = []
    for b in blocks:
        k = b[0]
        if k == "tbl":
            out += table(*b[1:])
        elif k == "b":
            out += [P("b", t) for t in b[1]]
        elif k == "h2" and annex:
            out.append(P("h3", b[1]))
        else:
            out.append(P(k, b[1]))
    return out

def insert_after(anchor, els):
    for e in els:
        anchor.addnext(e); anchor = e
    return anchor

def insert_before(anchor, els):
    for e in els:
        anchor.addprevious(e)

def set_text(p, new):
    """Replace a paragraph's text, keeping the first run's formatting."""
    ts = list(p.iter(q("t")))
    ts[0].text = new
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for t in ts[1:]: t.text = ""

def replace_text(p, old, new):
    full = text_of(p)
    assert old in full, (old, full[:80])
    for t in p.iter(q("t")):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new); return
    set_text(p, full.replace(old, new))

def clone_bullet(neigh, text):
    p = el("p"); p.append(copy.deepcopy(neigh.find(q("pPr")))); add_runs(p, text)
    return p

def next_table(after):
    k = after.getnext()
    while k is not None and k.tag != q("tbl"):
        k = k.getnext()
    return k

# ------------------------------------------------------------------ register existing captions
for p in top_paras():
    m = re.match(r"Table (\d+):", text_of(p).strip())
    if m:
        CAPTIONS["old" + m.group(1)] = p

# ------------------------------------------------------------------ 0. global text fixes and heading styles
for t in body.iter(q("t")):
    if t.text and "Maboko" in t.text:
        t.text = t.text.replace("Maboko", "Mabogo")
for p in top_paras():
    tx = text_of(p).strip()
    if style_of(p) == "Heading2" and tx == "Key Frameworks":
        p.find(q("pPr") + "/" + q("pStyle")).set(q("val"), "Heading3")
    if style_of(p) == "Heading3" and tx.startswith("6.6 "):
        p.find(q("pPr") + "/" + q("pStyle")).set(q("val"), "Heading2")
    if style_of(p) in ("Heading2", "Heading3"):           # remove stray italics in numbered H2s
        if re.match(r"\d+\.\d+ ", tx):
            for r in p.iter(q("rPr")):
                for tag in ("i", "iCs"):
                    for x in r.findall(q(tag)): r.remove(x)
    if style_of(p) == "Heading3" and tx.startswith("Harmonization Opportunities"):
        set_text(p, "Harmonisation Opportunities")

# ------------------------------------------------------------------ 1. Executive summary
harm = find_p(startswith("Harmonization: While national"))
insert_after(harm, [P("p", t) for _, t in C.EXEC_SUMMARY_ADD])
set_text(find_p(startswith("All findings are based on a desk review")), C.EXEC_SUMMARY_METHOD_NOTE)

# ------------------------------------------------------------------ 2. Chapter openers
for key, lead in C.LEADS.items():
    h = find_p(heading(key, 1))
    h.addnext(P("lead", lead))

# ------------------------------------------------------------------ 3. Chapter 1: MSME definition, method, ratings
meth_h = find_p(heading("Methodology and Limitations", 2))
insert_before(meth_h, render(C.SEC_1_4))
set_text(meth_h, "#.# Methodology, Data and Limitations")
old_p1 = find_p(startswith("The analysis is based on a structured desk review"), meth_h)
new_p1 = P("p", C.METHOD_P1)
old_p1.addprevious(new_p1)
first_bullet = old_p1.getnext()
old_p1.getparent().remove(old_p1)
anchor = insert_after(new_p1, [clone_bullet(first_bullet, g) for g in C.METHOD_GROUPS] + [P("p", C.METHOD_P2)])
lim = find_p(startswith("Limitations:"), meth_h)
lim = insert_after(lim, render(C.METHOD_LIMITS_ADD))
insert_after(lim, render(C.SEC_1_6))

# ------------------------------------------------------------------ 4. Chapter 2: constraints narrative, product terms, OKACOM, CORB
h23 = find_p(heading("Cross-Cutting Constraints", 2))
k = h23.getnext()
while k is not None and style_of(k) == "ListParagraph":
    nxt = k.getnext(); k.getparent().remove(k); k = nxt
anchor = insert_after(h23, render(C.SEC_2_3))
insert_after(anchor, render(C.SEC_2_4) + render(C.SEC_2_5) + render(C.SEC_2_6))
# OKACOM / CORB row in regional actors table
t2 = next_table(CAPTIONS["old2"]); rows = t2.findall(q("tr"))
newrow = copy.deepcopy(rows[-1])
for tc, txt in zip(newrow.findall(q("tc")), [
        "OKACOM / Cubango-Okavango River Basin (CORB) Fund",
        "Hybrid endowment and sinking fund for the basin states (Angola, Botswana, Namibia); USD 10m GEF grant approved May 2026",
        "Grant and first-loss layer for community livelihood, solar, water and eco-tourism investments in the basin portion of KAZA (sections 2.5–2.6)"]):
    ps = tc.findall(q("p"))
    for extra in ps[1:]: tc.remove(extra)
    set_text(ps[0], txt)
rows[-1].addnext(newrow)

# ------------------------------------------------------------------ 5. Country chapters
COUNTRY_TABLE = {"Angola": "old5", "Botswana": "old12", "Namibia": "old19", "Zambia": "old26", "Zimbabwe": "old33"}
for country, okey in COUNTRY_TABLE.items():
    tkey, cap, rows = C.TERMS[country]
    tb = next_table(CAPTIONS[okey])
    intro = P("p", "{T:%s} adds the indicative terms behind each rating in {T:%s}: typical pricing, tenor, ticket size and collateral, and the evidence on which the access rating is based (see section 1.6 for the scales). Where a figure is not published, the table says so." % (tkey, okey))
    insert_after(tb, [P("blank"), intro] + table(tkey, cap, C.TERMS_HEAD, rows, C.TERMS_W))

    ch = find_p(heading(country, 1))
    kf = find_p(lambda p: style_of(p) == "Heading3" and text_of(p).strip() == "Key Frameworks", ch)
    barriers = find_p(lambda p: style_of(p) == "Heading3" and text_of(p).strip() == "Barriers", kf)
    lastb = barriers.getprevious()
    while style_of(lastb) != "ListParagraph": lastb = lastb.getprevious()
    lastb.addnext(clone_bullet(lastb, C.REG_ADD[country]))
    if country in C.INFORMAL_BARRIER:
        en = find_p(lambda p: style_of(p) == "Heading3" and text_of(p).strip() == "Enablers", barriers)
        lb = en.getprevious()
        while style_of(lb) != "ListParagraph": lb = lb.getprevious()
        lb.addnext(clone_bullet(lb, C.INFORMAL_BARRIER[country]))
    prio = find_p(startswith(country + " priority"), ch)
    insert_after(prio, render(C.CLOSES[country]))

replace_text(find_p(startswith("Angola priority")), "grants and technical assistance for business formalization", "grants and technical assistance that support a formalisation pathway")
replace_text(find_p(lambda p: "monetary policy rates near 19" in text_of(p)),
             "monetary policy rates near 19–20%", "a central bank policy rate that has been cut from 18.5% at the end of 2025 to 14.75% in September 2026 (Forbes África Lusófona, 2026b)")
replace_text(find_p(lambda p: "14.5% in August 2025" in text_of(p)),
             "(CIEC Data, 2025)", "(CIEC Data, 2025), easing to 13.25% after the May 2026 cut (Lusaka Times, 2026)")
replace_text(find_p(lambda p: "Banking and Financial Services Act (2017)" in text_of(p)),
             "Banking and Financial Services Act (2017)", "Banking and Financial Services Act No. 9 of 2026 (replacing the 2017 Act)")
p_tff = find_p(lambda p: "RBZ Targeted Finance Facility" in text_of(p) and "on-lending cap" in text_of(p))
set_text(p_tff, "RBZ Targeted Finance Facility - At 12.5% since September 2026, with banks’ all-inclusive lending to productive sectors capped at 22.5% (The Herald, 2026), it could fund a tourism MSME window at partner banks.")

# ------------------------------------------------------------------ 6. New chapter 8 (communities)
synth = find_p(heading("Cross-Country Synthesis", 1))
insert_before(synth, render(C.CH8))

# ------------------------------------------------------------------ 7. Synthesis: scoring, barriers, harmonisation
t38 = next_table(CAPTIONS["old38"])
last = t38.findall(q("tr"))[-1]
vals = ["Overall readiness (score out of 5; section 9.2)", "Emerging (1.4)", "Ready with adjustments (3.4)", "Ready with adjustments (3.2)", "Ready to pilot (4.1)", "Ready to pilot (3.9)"]
for tc, v in zip(last.findall(q("tc")), vals):
    set_text(tc.find(q("p")), v)
after38 = t38.getnext()
insert_after(after38, render(C.SCORING))

bar_h = find_p(heading("Common Barriers", 2))
k = bar_h.getnext()
while k is not None and style_of(k) == "ListParagraph":
    nxt = k.getnext(); k.getparent().remove(k); k = nxt
insert_after(bar_h, render(C.BARRIERS))

harm_h = find_p(heading("harmonization Opportunities", 2))
set_text(harm_h, "#.# Harmonisation Opportunities at KAZA Level")
b1 = find_p(startswith("Common R-NBT eligibility"), harm_h); set_text(b1, "");
for r in list(b1.findall(q("r"))): b1.remove(r)
add_runs(b1, C.HARM_ELIGIBILITY)
b2 = find_p(startswith("Common definition of tourism MSME"), harm_h)
for r in list(b2.findall(q("r"))): b2.remove(r)
add_runs(b2, C.HARM_DEFINITION)
pay = find_p(startswith("Payments:"), harm_h)
insert_after(pay, [clone_bullet(pay, t) for t in C.HARM_ADD])

# ------------------------------------------------------------------ 8. Recommendations
t39 = next_table(CAPTIONS["old39"])
for tr in t39.findall(q("tr"))[1:]:
    tcs = tr.findall(q("tc")); first = text_of(tcs[0])
    if first.startswith("3"):
        set_text(tcs[1].find(q("p")), text_of(tcs[1]) + "; recognised groups and savings groups accepted as intermediaries")
        set_text(tcs[3].find(q("p")), text_of(tcs[3]) + "; CORB Fund in basin states")
    if first.startswith("5"):
        ps = tcs[1].findall(q("p"))
        for extra in ps[1:]: tcs[1].remove(extra)
        set_text(ps[0], "Grants and TA that help informal, community and micro-enterprises move along a formalisation pathway (starter tier accepting recognised groups and council licences), converting to loans or guarantees on milestones")
dp = find_p(startswith("Learn from low-uptake facilities"))
insert_after(dp, [clone_bullet(dp, t) for t in C.DESIGN_ADD])
ns_h = find_p(heading("Next Steps", 2))
insert_before(ns_h, render(C.SEC_10_3))
v1 = find_p(startswith("Validate institutional assessments"), ns_h)
v2 = v1.getnext()
for r in list(v2.findall(q("r"))): v2.remove(r)
add_runs(v2, C.NEXT_STEPS_FIX)
k = v2.getnext()
while k is not None and style_of(k) == "ListParagraph" and not text_of(k).strip():
    nxt = k.getnext(); k.getparent().remove(k); k = nxt
insert_after(v2, [clone_bullet(v2, t) for t in C.NEXT_STEPS_ADD])

# ------------------------------------------------------------------ 9. Acronyms (add + sort) and glossary
acr = doc.tables[0]._tbl
rows = acr.findall(q("tr")); template_row = rows[1]
existing = {text_of(r.findall(q("tc"))[0]).strip() for r in rows[1:]}
for a, m in C.ACRONYMS_ADD:
    if a in existing: continue
    nr = copy.deepcopy(template_row)
    for tc, v in zip(nr.findall(q("tc")), (a, m)):
        set_text(tc.find(q("p")), v)
    acr.append(nr)
rows = acr.findall(q("tr"))
for r in rows[1:]: acr.remove(r)
for r in sorted(rows[1:], key=lambda r: text_of(r.findall(q("tc"))[0]).strip().upper()):
    acr.append(r)

acr_title = find_p(startswith("Acronyms and Abbreviations"))
gl_title = copy.deepcopy(acr_title); set_text(gl_title, "Glossary of Terms")
gl_title.find(q("pPr")).insert(0, el("pageBreakBefore"))
gl = [gl_title, P("p", "Finance and conservation terms used in this report, in plain language.")]
gl += table("glossary", None, ["Term", "Meaning"], [[f"**{a}**", b] for a, b in C.GLOSSARY], [2600, 6426])
insert_after(acr, gl)

# ------------------------------------------------------------------ 10. References
ref_h = find_p(heading("References", 1))
if ref_h.find(q("pPr") + "/" + q("pageBreakBefore")) is None:
    ref_h.find(q("pPr")).insert(1, el("pageBreakBefore"))
refs = []
k = ref_h.getnext()
while k is not None and k.tag == q("p"):
    if text_of(k).strip() and not text_of(k).startswith("Web sources"):
        refs.append(k)
    k = k.getnext()
for r in refs[:]:
    if any(text_of(r).startswith(pfx) for pfx in C.REFS_REMOVE_PREFIX):
        r.getparent().remove(r); refs.remove(r)
tmpl = refs[0]
def ref_p(txt):
    p = el("p"); p.append(copy.deepcopy(tmpl.find(q("pPr")))); add_runs(p, txt, sz=20); return p
key = lambda s: re.sub(r"[^a-z0-9 ]", "", s.lower())
for txt in C.REFS_ADD:
    np_ = ref_p(txt)
    for r in refs:
        if key(text_of(r)) > key(txt):
            r.addprevious(np_); break
    else:
        refs[-1].addnext(np_)
    refs = sorted(refs + [np_], key=lambda e: list(body).index(e))
last_ref = refs[-1]

# ------------------------------------------------------------------ 11. Annex A (stakeholders) and Annex B (interview guide)
def table_rows(okey):
    t = next_table(CAPTIONS[okey])
    return [[text_of(tc).strip() for tc in tr.findall(q("tc"))] for tr in t.findall(q("tr"))[1:]]

stake = []
def addst(org, country, typ, role, interviewed="To be confirmed"):
    if org and org not in [s[0] for s in stake if s[1] == country]:
        stake.append([org, country, typ, role, interviewed])
for r in table_rows("old2"):
    addst(r[0], "Regional", "Regional financier / facility", r[2])
SETS = {"Angola": (4, 6, 7, 8), "Botswana": (11, 13, 14, 15), "Namibia": (18, 20, 21, 22), "Zambia": (25, 27, 28, 29), "Zimbabwe": (32, 34, 35, 36)}
for country, (bso, prv, fi, dfi) in SETS.items():
    for r in table_rows("old%d" % fi):
        addst(r[0], country, "Financial institution", r[4], "Yes (key informant interview)" if r[0].startswith("NASRIA") else "To be confirmed")
    for r in table_rows("old%d" % bso):
        addst(r[0], country, r[1] or "Stakeholder", r[2])
    for r in table_rows("old%d" % prv):
        addst(r[0], country, "Private sector (tourism)", r[2])
    for r in table_rows("old%d" % dfi):
        addst(r[0], country, "DFI / foundation / impact investor", r[2])
annexA = [P("h1", "Annex A: Stakeholder Database"),
          P("p", "This annex lists each stakeholder named in the report, with its country, type and role in MSME finance. The interview column records whether the organisation was consulted for this study; entries marked ‘To be confirmed’ should be completed from the interview log. Personal contact details are deliberately omitted so that the annex can be shared.")]
annexA += table("stakeholders", "Stakeholder database", ["No.", "Organisation", "Country", "Type", "Role in MSME finance", "Interviewed"],
                [[str(i + 1)] + s for i, s in enumerate(stake)], [450, 2250, 950, 1450, 2726, 1200])
insert_after(insert_after(last_ref, annexA), render(C.ANNEX_B, annex=True))

# ------------------------------------------------------------------ 12. Renumber headings
h1n = 0; h2n = 0
for p in top_paras():
    st, tx = style_of(p), text_of(p).strip()
    if st == "Heading1":
        m = re.match(r"(#|\d+)\.\s+(.*)", tx)
        if m:
            h1n += 1; h2n = 0; set_text_keep = "%d. %s" % (h1n, m.group(2))
            set_text(p, set_text_keep)
        else:
            h2n = 0
    elif st == "Heading2":
        m = re.match(r"(#\.#|\d+\.\d+)\s+(.*)", tx)
        if m:
            h2n += 1; set_text(p, "%d.%d %s" % (h1n, h2n, m.group(2)))

# ------------------------------------------------------------------ 13. Renumber tables and resolve {T:key}
num = {}
order = [p for p in top_paras() if re.match(r"Table \d+:", text_of(p).strip())]
inv = {id(v): k for k, v in CAPTIONS.items()}
for i, p in enumerate(order, 1):
    k = inv.get(id(p))
    if k: num[k] = i
    ts = [t for t in p.iter(q("t")) if t.text]
    ts[0].text = re.sub(r"Table \d+:", "Table %d:" % i, ts[0].text, count=1)
for t in body.iter(q("t")):
    if t.text and "{T:" in t.text:
        t.text = re.sub(r"\{T:(\w+)\}", lambda m: "Table %d" % num[m.group(1)], t.text)
# the only V4 body cross-reference ("Table 1 summarizes") stays valid: Table 1 is unchanged
assert num["old1"] == 1

# ------------------------------------------------------------------ 14. Rebuild the TOC
sdt = next(k for k in body if k.tag == q("sdt"))
sc = sdt.find(q("sdtContent")); kids = list(sc)
entries = [k for k in kids if k.find(q("hyperlink")) is not None]
t1 = copy.deepcopy(next(e for e in entries[1:] if style_of(e) == "TOC1"))
t2 = copy.deepcopy(next(e for e in entries if style_of(e) == "TOC2"))
head_runs = [copy.deepcopy(r) for r in entries[0] if r.tag == q("r")]
for e in entries: sc.remove(e)
new_entries = []
for p in top_paras():
    st = style_of(p)
    if st not in ("Heading1", "Heading2"): continue
    tx = text_of(p).strip()
    bs = p.find(q("bookmarkStart"))
    if bs is None or not bs.get(q("name"), "").startswith("_Toc"):
        name = bookmark_wrap(p)
    else:
        name = bs.get(q("name"))
    e = copy.deepcopy(t1 if st == "Heading1" else t2)
    hl = e.find(q("hyperlink")); hl.set(q("anchor"), name)
    ts = list(hl.iter(q("t")))
    ts[0].text = tx
    ts[-1].text = str(pages.get(tx, 0))
    for it in hl.iter(q("instrText")):
        it.text = " PAGEREF %s \\h " % name
    new_entries.append(e)
first = new_entries[0]
for i, r in enumerate(head_runs, 1):   # field begin, instruction, separate go right after pPr
    first.insert(i, r)
for e in reversed(new_entries):
    sc.insert(0, e)

# keep the empty section-break paragraph after the TOC from spilling onto its own page
secp = kids[-1]
spp = secp.find(q("pPr"))
for x in spp.findall(q("spacing")) + spp.findall(q("rPr")): spp.remove(x)
spp.insert(0, el("spacing", before=0, after=0, line=20, lineRule="exact"))
rp = el("rPr"); sub(rp, "sz", val=2); sub(rp, "szCs", val=2)
sp_sect = spp.find(q("sectPr"))
(sp_sect.addprevious(rp) if sp_sect is not None else spp.append(rp))

# ask Word to refresh fields (TOC page numbers) on open
settings = doc.settings.element
if settings.find(q("updateFields")) is None:
    uf = el("updateFields", val="true")
    for tag in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids"):
        nxt = settings.find(q(tag))
        if nxt is not None:
            nxt.addprevious(uf); break
    else:
        settings.append(uf)

doc.save(out)
json.dump({"headings": [text_of(p).strip() for p in top_paras() if style_of(p) in ("Heading1", "Heading2")],
           "tables": len(order), "stakeholders": len(stake)}, open(out + ".info.json", "w"), indent=1, ensure_ascii=False)
print("saved", out, "tables", len(order), "stakeholders", len(stake))
